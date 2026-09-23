"""
DAL del modulo financiero (Dashboard Unit Economics).

REGLA DE ORO: este modulo NO usa SQL Server. No importa app.database ni pyodbc.
La fuente de datos es un Excel que persiste en BE/datos/ y se puede reemplazar
en caliente desde la vista "Analisis Financiero" de VISOR360.

Responsabilidades:
- Resolver la ruta del Excel de forma portable (nunca absoluta de una maquina).
- Validar que un Excel candidato traiga las 12 hojas que el tablero necesita.
- Reemplazar el Excel vigente con respaldo con fecha.
- Limpiar la cache en memoria para que las vistas se recalculen tras una subida.
"""
from __future__ import annotations

import gc
import logging
import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from app.financiero import data_loader as dl

logger = logging.getLogger(__name__)

# Hojas que el tablero necesita. Portadas tal cual de la fuente (dashboard_api.py
# y app_pages/actualizar_datos.py del proyecto original).
HOJAS_REQUERIDAS = {
    "P&L Consolidated", "P&L Servicing", "P&L Master", "P&L NPL",
    "7Actuals info", "7Actuals info Peru",
    "Monthly 12 Budget wo Peru", "Monthly 12 Budget Peru",
    "Budget ALL SPV BU Detailed", "Fee Services", "Cost per Million",
    "Profitability Analysis",
}


def ruta_excel() -> Path:
    """Ruta del Excel vigente, resuelta de forma portable (ver data_loader).

    Se recalcula en cada llamada a proposito: si un usuario sube un Excel nuevo,
    la resolucion (que elige el .xlsx mas reciente de BE/datos/) lo toma solo.
    """
    return dl._resolver_ruta_excel()


def info_excel() -> dict:
    """Metadatos del Excel vigente para mostrarlos en la UI."""
    ruta = ruta_excel()
    if not ruta.exists():
        return {"existe": False, "nombre": ruta.name, "ubicacion": str(ruta.parent)}
    return {
        "existe": True,
        "nombre": ruta.name,
        "ubicacion": str(ruta.parent),
        "tamano_mb": round(ruta.stat().st_size / 1_048_576, 2),
        "actualizado": datetime.fromtimestamp(ruta.stat().st_mtime).strftime("%Y-%m-%d %H:%M"),
    }


def validar_hojas(ruta: Path) -> list[str]:
    """Devuelve la lista de hojas requeridas que faltan en el archivo dado.

    Lista vacia == archivo valido. Lanza si el archivo no es un Excel legible.
    """
    wb = load_workbook(ruta, read_only=True)
    try:
        presentes = set(wb.sheetnames)
    finally:
        wb.close()
    # En Windows, read_only mantiene el handle del .xlsx hasta que el objeto
    # Workbook -y su archivo interno- salen del scope. Soltar la referencia y
    # forzar un ciclo del GC aqui (no en el borrado) evita el WinError 32 que
    # antes dejaba residuos tablero_upload_*.xlsx en %TEMP%.
    del wb
    gc.collect()
    return sorted(HOJAS_REQUERIDAS - presentes)


def validar_hojas_bytes(contenido: bytes) -> list[str]:
    """Valida las 12 hojas SIN escribir nada a disco (workbook desde memoria).

    Evita por completo el temporal tablero_upload_*.xlsx: en Windows, un .xlsx
    en disco recien escrito queda con el handle retenido (read_only + antivirus)
    y el borrado fallaba dejando residuos en %TEMP%. Con BytesIO no hay archivo
    temporal que borrar. Lanza si el contenido no es un Excel legible.
    """
    import io
    wb = load_workbook(io.BytesIO(contenido), read_only=True)
    try:
        presentes = set(wb.sheetnames)
    finally:
        wb.close()
    del wb
    gc.collect()
    return sorted(HOJAS_REQUERIDAS - presentes)


def reemplazar_excel(contenido: bytes) -> dict:
    """Reemplaza el Excel vigente por el contenido recibido (bytes).

    Flujo seguro:
    1. Valida las 12 hojas en MEMORIA (BytesIO), sin temporal en disco.
    2. Si falta alguna hoja, aborta sin tocar el archivo vigente.
    3. Respalda el vigente como {stem}_respaldo_{YYYYMMDD_HHMM}{suffix}.
    4. Escribe el nuevo (write_bytes: en Windows move a destino existente puede
       fallar o dejar huerfanos).
    5. Limpia la cache en memoria para que las vistas se recalculen.

    Devuelve {"ok": bool, "backup": str|None, "faltantes": [...], "error": str|None}.
    """
    ruta = ruta_excel()
    ruta.parent.mkdir(parents=True, exist_ok=True)
    try:
        # Validacion en memoria: no se crea ningun .xlsx temporal, asi que no
        # hay handle que Windows retenga ni residuo que limpiar en %TEMP%.
        faltantes = validar_hojas_bytes(contenido)
        if faltantes:
            return {"ok": False, "backup": None, "faltantes": faltantes, "error": None}

        marca = (
            datetime.fromtimestamp(ruta.stat().st_mtime).strftime("%Y%m%d_%H%M")
            if ruta.exists() else "sin_anterior"
        )
        backup = ruta.with_name(f"{ruta.stem}_respaldo_{marca}{ruta.suffix}")
        if ruta.exists():
            shutil.copy2(ruta, backup)

        ruta.write_bytes(contenido)
        limpiar_cache()
        logger.info("Excel reemplazado: %s (respaldo: %s)", ruta.name, backup.name)
        return {"ok": True, "backup": backup.name, "faltantes": [], "error": None}
    except Exception as exc:  # noqa: BLE001 -- se reporta al cliente tal cual
        logger.exception("Fallo al reemplazar el Excel")
        return {"ok": False, "backup": None, "faltantes": [], "error": str(exc)}


def limpiar_cache() -> None:
    """Limpia la cache de las funciones de carga de data_loader.

    data_loader usa functools.lru_cache (antes @st.cache_data). lru_cache expone
    .cache_clear() en cada funcion decorada; se recorren todas las del modulo.
    """
    limpiadas = 0
    for nombre in dir(dl):
        funcion = getattr(dl, nombre)
        if callable(funcion) and hasattr(funcion, "cache_clear"):
            try:
                funcion.cache_clear()
                limpiadas += 1
            except Exception:  # noqa: BLE001
                pass
    logger.info("Cache de data_loader limpiada (%s funciones)", limpiadas)
